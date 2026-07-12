from pathlib import Path
import contextlib
import sqlite3
from datetime import datetime
import json
import shutil
from collections import defaultdict
import re
import secrets
import pandas as pd
from fresfolio.utils import tools

extras_installed = tools.has_extras()
if extras_installed:
    from omilayers import Omilayers
    import duckdb
    from docx import Document
    import pymupdf4llm
    pymupdf4llm.use_layout(False)

APPDIR = Path("~/fresfolio").expanduser()
APPDB = APPDIR.joinpath("fresfolio.db")

if APPDIR.exists():
    from fresfolio.renderers.html_renderer import HtmlRenderer, PDFRenderer

class AppINIT:

    def __init__(self) -> None:
        print(">>> Initializing fresfolio app...")
        if not self.__check_init_conditions_passed:
            exit(">>> Fresfolio was not initialized successfully.")
        print(">>> fresfolio initialized successfully.")
        self.projectsDir = Path(tools.get_app_setting("projectsDir")).expanduser()

    @property
    def __check_init_conditions_passed(self) -> bool:
        conds = [
            self.__app_dir_initialized,
            self.__app_db_initialized,
            self.__projects_dir_initialized
        ]

        if sum(conds) != len(conds):
            return False
        return True

    @property
    def __app_dir_initialized(self) -> bool:
        if not APPDIR.exists():
            try:
                APPDIR.mkdir(exist_ok=False)
            except Exception:
                print("[ERROR] Cannot initialize Fresfolio app directory.")
                tools.log_traceback()
                return False
            print("[OK] app directory created.")
            return True
        else:
            print("[OK] app directory exists.")
            return True

    @property
    def __app_db_initialized(self) -> bool:
        if not APPDB.exists():
            try:
                with contextlib.closing(sqlite3.connect(APPDB)) as conn:
                    with contextlib.closing(conn.cursor()) as c:
                        query = """
                        CREATE TABLE settings(
                        key TEXT,
                        value TEXT
                        )
                        """
                        c.execute(query)

                        query = """
                        CREATE TABLE projects(
                        id INTEGER PRIMARY KEY,
                        uuid TEXT,
                        name TEXT,
                        path TEXT,
                        description TEXT,
                        started TEXT,
                        finished TEXT
                        )
                        """
                        c.execute(query)
                        conn.commit()

                        query = """
                        INSERT INTO settings 
                        (key,value) 
                        VALUES (?,?)
                        """
                        c.execute(query, ("projectsDir", str(APPDIR.joinpath("projects"))))
                        c.execute(query, ("secret_key", secrets.token_urlsafe(32)))
                        c.execute(query, ("has_set_uuids", 1))

                        conn.commit()
            except Exception:
                print("[ERROR] Cannot initialize app database.")
                tools.log_traceback()
                return False
            print("[OK] app database created.")
            return True
        else:
            print("[OK] app database exists.")
            return True

    @property
    def __projects_dir_initialized(self) -> bool:
        try:
            projectsDir = tools.get_app_setting("projectsDir")
        except Exception:
            print("[ERROR] Cannot get setting for projects directory.")
            tools.log_traceback()
            return False

        if projectsDir is None:
            print("[ERROR] projects directory is None.")
            return False

        projectsDir = Path(projectsDir).expanduser()
        if not projectsDir.exists():
            try:
                projectsDir.mkdir(exist_ok=False)
            except Exception:
                print("[ERROR] Cannot create projects directory.")
                tools.log_traceback()
                return False
            print("[OK] projects directory created.")
            return True
        else:
            print("[OK] projects directory exists.")
            return True

class ProjectsUtils:

    def __init__(self):
        self.projectsDir = tools.get_app_setting("projectsDir")

    def get_projects(self) -> list:
        with contextlib.closing(sqlite3.connect(APPDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = "SELECT uuid,name,description,started FROM projects"
                c.execute(query)
                rows = c.fetchall()
        if rows:
            return [{"id":row[0], "name": row[1], "description":row[2], "started":row[3]} for row in rows]
        return []

    def project_exists(self, projectName:str) -> bool:
        projectName = projectName
        with contextlib.closing(sqlite3.connect(APPDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = "SELECT 1 FROM projects WHERE name=(?) LIMIT 1"
                c.execute(query, (projectName,))
                row = c.fetchone()
        if row:
            return True
        return False

    def project_is_created(self, projectName:str, projectDescription:str) -> bool:
        today = datetime.today().strftime('%Y-%m-%d')
        projectDirectory = Path(self.projectsDir).joinpath(projectName)
        try:
            projectDirectory.mkdir(exist_ok=False)
            projectDirectory.joinpath("sections").mkdir(exist_ok=False)
        except Exception:
            tools.log_traceback()
            return False

        projectDB = str(projectDirectory.joinpath("project.db"))
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    CREATE TABLE notebooks(
                    id INTEGER PRIMARY KEY,
                    notebook TEXT,
                    date TEXT
                    )
                    """
                    c.execute(query)

                    query = """
                    CREATE TABLE chapters(
                    id INTEGER PRIMARY KEY,
                    chapter TEXT,
                    notebookID INTEGER,
                    date TEXT
                    )
                    """
                    c.execute(query)

                    query = """
                    CREATE TABLE sections(
                    id INTEGER PRIMARY KEY,
                    section TEXT,
                    tags TEXT,
                    content TEXT,
                    date TEXT
                    )
                    """
                    c.execute(query)

                    query = """
                    CREATE TABLE chapters_sections_links(
                    id INTEGER PRIMARY KEY,
                    chapterID INTEGER,
                    sectionID INTEGER
                    )
                    """
                    c.execute(query)
                    conn.commit()
        except Exception:
            tools.log_traceback()
            if projectDirectory.exists():
                shutil.rmtree(projectDirectory)
            return False

        try:
            with contextlib.closing(sqlite3.connect(APPDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    INSERT INTO projects (uuid,name,path,description,started) 
                    VALUES (?,?,?,?,?)
                    """
                    c.execute(query, (tools.generate_uuid(), projectName, str(projectDirectory), projectDescription, today))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            if projectDirectory.exists():
                shutil.rmtree(projectDirectory)
            return False
        return True

    def project_is_imported(self, projectName:str, projectPath:str) -> bool:
        today = datetime.today().strftime('%Y-%m-%d')
        projectDescription = "Imported project. No available description."
        try:
            with contextlib.closing(sqlite3.connect(APPDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    INSERT INTO projects (uuid,name,path,description,started) 
                    VALUES (?,?,?,?,?)
                    """
                    c.execute(query, (tools.generate_uuid(), projectName, projectPath, projectDescription, today))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def project_path_is_set(self, projectName:str, projectPath:str) -> bool:
        try:
            with contextlib.closing(sqlite3.connect(APPDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "UPDATE projects SET path=(?) WHERE name=(?)"
                    c.execute(query, (str(projectPath), projectName))
                    conn.commit()
        except Exception: 
            tools.log_traceback()
            return False
        return True

    def get_notebooks_for_project(self, projectID:str) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = "SELECT id, notebook, date from notebooks"
                c.execute(query)
                notebooks =  c.fetchall()
        return notebooks

    @staticmethod
    def get_chapters_for_notebook_of_project(projectID:str, notebookID:int) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = "SELECT id,chapter,date FROM chapters WHERE notebookID=(?)"
                c.execute(query, (notebookID,))
                chapters = c.fetchall()
        return chapters

    @staticmethod
    def get_chapters_ids_for_notebook(projectID:str, notebookID:int) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = "SELECT id FROM chapters WHERE notebookID=(?)"
                c.execute(query, (notebookID,))
                chapters = c.fetchall()
        if not chapters:
            return []
        return [c[0] for c in chapters]

    def get_notebooks_and_chapters_for_project(self, projectID:str) -> dict: 
        notebooks = self.get_notebooks_for_project(projectID)
        sidebarDataList = []
        if notebooks:
            # sort notebooks by name
            notebooks = sorted(notebooks, key=lambda field: field[1])
            for notebook in notebooks:
                notebookID, notebookName, notebookDate = notebook
                JSON = {"notebookID":notebookID, "notebookName":notebookName, "notebookDate":notebookDate, "chapters":[]}
                chapters = self.get_chapters_for_notebook_of_project(projectID, notebookID)
                if chapters:
                    # sort chapters by name
                    chapters = sorted(chapters, key=lambda field: field[1])
                    for chapter in chapters:
                        chapterID, chapterName, chapterDate = chapter 
                        JSON['chapters'].append({
                            "chapterID": chapterID,
                            "chapterName": chapterName,
                            "chapterDate": chapterDate
                            })
                sidebarDataList.append(JSON)
        return sidebarDataList

    def notebook_exists(self, projectID:str, notebookName:str) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = 'SELECT EXISTS(SELECT 1 FROM notebooks WHERE notebook=(?))'
                c.execute(query, (notebookName, ))
                result = c.fetchone()[0]
        if result == 1:
            return True
        return False

    def notebook_is_created(self, projectID:str, notebookName:str) -> bool:
        today = datetime.today().strftime('%Y-%m-%d')
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    INSERT INTO notebooks (notebook,date) 
                    VALUES (?,?)
                    """
                    c.execute(query, (notebookName, today))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def notebook_name_is_set(self, projectID:str, notebookID:int, newNotebookName:str) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    UPDATE notebooks 
                    SET notebook=(?) 
                    WHERE id=(?)
                    """
                    c.execute(query, (newNotebookName,notebookID))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def project_description_is_set(self, projectID:str, newProjectDescription:str) -> bool:
        try:
            with contextlib.closing(sqlite3.connect(APPDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    UPDATE projects 
                    SET description=(?) 
                    WHERE uuid=(?)
                    """
                    c.execute(query, (newProjectDescription, projectID))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def project_name_is_set(self, projectID:str, newProjectName:str) -> bool:
        try:
            oldProjectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            if not Path(oldProjectDirectory).exists():
                return False
            newProjectName = newProjectName
            newProjectDirectory = Path(oldProjectDirectory).parent.joinpath(newProjectName)
            Path(oldProjectDirectory).rename(newProjectDirectory)

            if not newProjectDirectory.exists():
                return False
        
            with contextlib.closing(sqlite3.connect(APPDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    UPDATE projects 
                    SET name=(?), path=(?)
                    WHERE uuid=(?)
                    """
                    c.execute(query, (newProjectName, str(newProjectDirectory), projectID))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def new_section_directory_created(self, projectID:str, sectionID:str) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            Path(projectDirectory).joinpath(f"sections/{sectionID}").mkdir(exist_ok=False)
        except Exception:
            tools.log_traceback()
            return False
        return True

    def section_in_db_exists(self, projectID:str, sectionID:int) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = 'SELECT EXISTS(SELECT 1 FROM sections WHERE id=(?))'
                c.execute(query, (sectionID,))
                result = c.fetchone()[0]
        if result == 1:
            return True
        return False

    def section_directory_exists(self, projectID:str, sectionID:int) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        sectionDirPath = Path(projectDirectory).joinpath(f"sections/{sectionID}")
        if sectionDirPath.exists():
            return True
        return False

    def section_directory_is_deleted(self, projectID:str, sectionID:int) -> bool:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            sectionDirPath = Path(projectDirectory).joinpath(f"sections/{sectionID}")
            shutil.rmtree(str(sectionDirPath))
        except Exception:
            tools.log_traceback()
            return False
        return True

    def section_in_db_is_deleted(self, projectID:str, sectionID:int) -> bool:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "DELETE FROM sections WHERE id=(?)"
                    c.execute(query, (sectionID,))

                    query = "DELETE FROM chapters_sections_links WHERE sectionID=(?)"
                    c.execute(query, (sectionID,))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def chapter_exists(self, projectID:str, notebookID:int, chapterName:str) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = 'SELECT EXISTS(SELECT 1 FROM chapters WHERE notebookID=(?) AND chapter=(?))'
                c.execute(query, (notebookID, chapterName))
                result = c.fetchone()[0]
        if result == 1:
            return True
        return False

    def chapter_is_created(self, projectID:str, notebookID:int, chapterName:str) -> bool: 
        today = datetime.today().strftime('%Y-%m-%d')
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    INSERT INTO chapters (chapter,notebookID,date) 
                    VALUES (?,?,?)
                    """
                    c.execute(query, (chapterName, notebookID, today))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def chapter_name_is_set(self, projectID:str, chapterID:int, newChapterName:str) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    UPDATE chapters
                    SET chapter=(?) 
                    WHERE id=(?)
                    """
                    c.execute(query, (newChapterName, chapterID))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def get_sections_IDs_for_chapter(self, projectID:str, chapterID:int) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "SELECT sectionID FROM chapters_sections_links WHERE chapterID=(?)"
                    c.execute(query, (chapterID,))
                    sectionsIDs = c.fetchall()
            if sectionsIDs:
                return [s[0] for s in sectionsIDs]
            return []
        except Exception:
            tools.log_traceback()
            return []

    def get_chapter_sections(self, projectID:str, chapterID:int) -> list:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            sectionsIDs = self.get_sections_IDs_for_chapter(projectID, chapterID)
            sectionsResults = [] 
            if sectionsIDs:
                cols = ['id', 'section', 'tags', 'content', 'date']
                with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                    with contextlib.closing(conn.cursor()) as c:
                        query = """
                        SELECT {0} FROM sections 
                        WHERE id IN ({1})
                        """.format(','.join(cols), ', '.join('?' for _ in sectionsIDs))
                        c.execute(query, sectionsIDs)
                        sectionsResults = c.fetchall()
            sections = []
            if sectionsResults:
                colsMapper = {
                            "id":"ID",
                            "section":"title",
                            "tags":"tags",
                            "content":"content",
                            "date":"sectionDate"
                                }
                sectionsJSON = {} # A way to retain the order of sections in sectionsIDs.
                for result in sectionsResults:
                    kwargs = {colsMapper[col]:value for col,value in zip(cols,result)}
                    if not kwargs['content'].strip("\n"):
                        kwargs['content'] = "Section content is emtpy."
                    kwargs['projectID'] = projectID
                    kwargs['projectName'] = tools.get_project_name_based_on_id(projectID)
                    kwargs['section_dir_exists'] = int(Path(projectDirectory).joinpath(f"sections/{kwargs['ID']}").exists())
                    section = SectionUtils(**kwargs)
                    sectionsJSON[kwargs['ID']] = section.render_content_to_html()
                sections = [sectionsJSON[ID] for ID in sectionsIDs]
            return sections
        except Exception:
            tools.log_traceback()
            return []

    def get_sections_based_on_tag(self, projectID:str, tag:str, descending=False) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        sections = [] 
        try:
            cols = ['id', 'section', 'tags', 'content', 'date']
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    if descending:
                        query = """
                        SELECT {0} FROM sections 
                        WHERE tags LIKE "%{1}%" 
                        ORDER BY id DESC
                        """.format(','.join(cols), tag)
                    else:
                        query = """
                        SELECT {0} FROM sections 
                        WHERE tags LIKE "%{1}%" 
                        ORDER BY id
                        """.format(','.join(cols), tag)
                    c.execute(query)
                    sections = c.fetchall()
        except Exception:
            tools.log_traceback()
            return sections
        return sections

    def get_chat_sections_content(self, projectID:str) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        sections = [] 
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    SELECT content FROM sections 
                    WHERE tags LIKE "%ai-chat%" 
                    ORDER BY id
                    """
                    c.execute(query)
                    sections = c.fetchall()
        except Exception:
            tools.log_traceback()
            return sections
        return [s[0] for s in sections]


    def get_section_content_rendered(self, projectID:str, sectionID:int, render_type:str='html') -> list:
        """Render section content to HTML"""
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        cols = ['id', 'section', 'tags', 'content', 'date']
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = f"""
                SELECT {','.join(cols)} FROM sections 
                WHERE id=(?)
                """
                c.execute(query, (sectionID,))
                result = c.fetchone()
        if result:
            colsMapper = {
                        "id":"ID",
                        "section":"title",
                        "tags":"tags",
                        "content":"content",
                        "date":"sectionDate"
                            }
            kwargs = {colsMapper[col]:value for col,value in zip(cols,result)}
            if not kwargs['content'].strip("\n"):
                kwargs['content'] = "Section content is emtpy."
            kwargs['projectID'] = projectID
            kwargs['projectName'] = tools.get_project_name_based_on_id(projectID)
            kwargs['section_dir_exists'] = int(Path(projectDirectory).joinpath(f"sections/{kwargs['ID']}").exists())
            section = SectionUtils(**kwargs)
            if render_type == "html":
                return section.render_content_to_html()
            elif render_type == "pdf":
                return section.render_content_to_pdf()
        return {}

    def get_section_raw_content(self, projectID:str, sectionID:int) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = """
                SELECT content FROM sections 
                WHERE id=(?)
                """
                c.execute(query, (sectionID,))
                result = c.fetchone()
        if result:
            return result[0]
        return ""

    def get_section_title(self, projectID:str, sectionID:int) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = """
                SELECT section FROM sections 
                WHERE id=(?)
                """
                c.execute(query, (sectionID,))
                result = c.fetchone()
        if result:
            return result[0]
        return ""

    def create_section_in_db(self, projectID:str, chapterID:int) -> int:
        today = datetime.today().strftime('%Y-%m-%d')
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        sectionID = None
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = """
                INSERT INTO sections 
                (section, tags, content, date)
                VALUES (?,?,?,?)
                """
                c.execute(query, ('New section', json.dumps([]), '', today))
                sectionID = c.lastrowid
                conn.commit()

                query = """
                INSERT INTO chapters_sections_links 
                (chapterID,sectionID) 
                VALUES (?,?)
                """
                c.execute(query, (chapterID, sectionID))
                conn.commit()
        return sectionID

    def insert_section_in_db(self, projectID:str, title:str, content:str, tags:list) -> int:
        today = datetime.today().strftime('%Y-%m-%d')
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        sectionID = None
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = """
                    INSERT INTO sections 
                    (section, tags, content, date)
                    VALUES (?,?,?,?)
                    """
                    c.execute(query, (title, json.dumps(tags), content, today))
                    sectionID = c.lastrowid
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return sectionID
        return sectionID

    def section_title_is_set(self, projectID:str, sectionID:int, newSectionTitle:str) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "UPDATE sections SET section=(?) WHERE id=(?)"
                    c.execute(query, (newSectionTitle, sectionID))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def section_content_is_set(self, projectID:str, sectionID:int, newSectionContent:str) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "UPDATE sections SET content=(?) WHERE id=(?)"
                    c.execute(query, (newSectionContent, sectionID))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def section_tags_is_set(self, projectID:str, sectionID:int, sectionTags:list) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "UPDATE sections SET tags=(?) WHERE id=(?)"
                    c.execute(query, (json.dumps(sectionTags), sectionID))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def get_section_tags(self, projectID:str, sectionID:int) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "SELECT tags FROM sections WHERE id=(?)"
                    c.execute(query, (sectionID, ))
                    result = c.fetchone()
            tags = []
            if result:
                tags = json.loads(result[0])
            return tags
        except Exception:
            tools.log_traceback()
            return []

    def check_which_section_IDs_exist_in_db(self, projectID:str, sectionsIDs:list) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        placeholders = ','.join('?' for _ in sectionsIDs)
        with contextlib.closing(sqlite3.connect(projectDB)) as conn:
            with contextlib.closing(conn.cursor()) as c:
                query = f"SELECT id FROM sections WHERE id IN ({placeholders})"
                c.execute(query, sectionsIDs)
                results = c.fetchall()
        if results:
            return [res[0] for res in results]
        return []

    def section_directory_is_created(self, projectID:str, sectionID:int) -> tuple:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        sectionDirectoryPath = Path(projectDirectory).joinpath(f"sections/{sectionID}")
        if sectionDirectoryPath.exists():
            return ("Section directory exists", False) 
        try:
            sectionDirectoryPath.mkdir(exist_ok=False)
        except Exception:
            tools.log_traceback()
            return ("Cannot create section directory", False)
        return ("", True)

    def get_sections_IDs_based_on_search_bar_query(self, projectID:str, queryTerms:str) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        projects = {}
        query = "SELECT id FROM sections WHERE"
        if queryTerms.startswith("id:"):
            try:
                IDs = queryTerms.split(":")[-1]
                if "," in IDs:
                    IDs = [int(x.strip(" ")) for x in IDs.split(",") if x]
                else:
                    IDs = [int(IDs.strip(" "))]
            except Exception:
                tools.log_traceback()
                IDs = []
            projects[projectID] = IDs
        else:
            if "," in queryTerms:
                connector = "OR"
                queryTerms = queryTerms.replace(",", "")
            else:
                connector = "AND"

            if queryTerms.strip(" ")[1] != ":":
                queryTerms = "a:" + queryTerms

            pattern = r'([a-z]):\s*(.*?)(?=(?: [a-z]:|$))'
            queryTags = defaultdict(list)
            for key, value_str in re.findall(pattern, queryTerms):
                # Extract words or quoted phrases
                words = re.findall(r'"[^"]+"|\S+', value_str)
                # Clean quotes and trailing punctuation like commas
                cleaned_words = [w.strip('"').strip(',') for w in words]
                queryTags[key].extend(cleaned_words)

            allTerms  = list()
            sectionTerms =  list()
            tagsTerms = list()
            contentTerms = list()
            _projects = list()

            if queryTags['p']:
                if "all" in queryTags['p']:
                    allProjects = tools.get_projects_names_and_paths()
                    for item in allProjects:
                        projectName, projectPath = item
                        _projectID = tools.get_project_ID_based_on_name(projectName)
                        _projectDirectory, _projectDB = tools.get_paths_for_project_dir_and_db(_projectID)
                        _projects.append((_projectID, _projectDB))
                else:
                    try:
                        for projectName in queryTags['p']:
                            _projectID = tools.get_project_ID_based_on_name(projectName)
                            if _projectID is None:
                                raise ValueError(f"Project '{projectName}' does not exist.")
                                return []
                            _projectDirectory, _projectDB = tools.get_paths_for_project_dir_and_db(_projectID)
                            _projects.append((_projectID, _projectDB))
                    except Exception:
                        tools.log_traceback()
                        return []
            else:
                _projects.append((projectID, projectDB))

            if queryTags['a']:
                for term in queryTags['a']:
                    allTerms.append('(tags LIKE "%{0}%" OR section LIKE "%{0}%" OR content LIKE "%{0}%")'.format(term))

            if queryTags['s']:
                for term in queryTags['s']:
                    sectionTerms.append('section LIKE "%{0}%"'.format(term))

            if queryTags['t']:
                for term in queryTags['t']:
                    tagsTerms.append('tags LIKE "%{0}%"'.format(term))

            if queryTags['c']:
                for term in queryTags['c']:
                    contentTerms.append('content LIKE "%{0}%"'.format(term))

            if allTerms:
                if query.endswith("WHERE"):
                    query += " "
                else:
                    query += f" {connector} "
                query += f" {connector} ".join(allTerms)
            if sectionTerms:
                if query.endswith("WHERE"):
                    query += " "
                else:
                    query += f" {connector} "
                query += f" {connector} ".join(sectionTerms)
            if tagsTerms:
                if query.endswith("WHERE"):
                    query += " "
                else:
                    query += f" {connector} "
                query += f" {connector} ".join(tagsTerms)
            if contentTerms:
                if query.endswith("WHERE"):
                    query += " "
                else:
                    query += f" {connector} "
                query += f" {connector} ".join(contentTerms)

            for item in _projects:
                _projectID, _projectDB = item
                with contextlib.closing(sqlite3.connect(_projectDB)) as conn:
                    with contextlib.closing(conn.cursor()) as c:
                        c.execute(query)
                        IDs = c.fetchall()
                if IDs:
                    IDs = [res[0] for res in IDs]
                projects[_projectID] = IDs
        return projects

    def notebook_is_deleted(self, projectID:str, notebookID:int, keep_sections:bool) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            sectionsIDs = []
            if not keep_sections:
                chapters = ProjectsUtils.get_chapters_for_notebook_of_project(projectID, notebookID)
                if chapters:
                    chaptersIDs = [chapter[0] for chapter in chapters]
                    sectionsIDs = []
                    for chapterID in chaptersIDs:
                        sectionsIDs.extend(self.get_sections_IDs_for_chapter(projectID, chapterID))
                    if sectionsIDs:
                        sectionsIDs = list(set(sectionsIDs))

            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    if sectionsIDs:
                        placeholders = ','.join('?' for _ in sectionsIDs)
                        query = f"DELETE FROM sections WHERE id IN ({placeholders})"
                        c.execute(query, sectionsIDs)

                        query = f"DELETE FROM chapters_sections_links WHERE sectionID IN ({placeholders})"
                        c.execute(query, sectionsIDs)

                    query = "DELETE FROM chapters WHERE notebookID=(?)"
                    c.execute(query, (notebookID,))

                    query = "DELETE FROM notebooks WHERE id=(?)"
                    c.execute(query, (notebookID,))
                    conn.commit()
            return True
        except Exception:
            tools.log_traceback()
            return False

    def chapter_is_deleted(self, projectID:str, chapterID:int, keep_sections:bool) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            sectionsIDs = []
            if not keep_sections:
                sectionsIDs = self.get_sections_IDs_for_chapter(projectID, chapterID)
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    if sectionsIDs:
                        placeholders = ','.join('?' for _ in sectionsIDs)
                        query = f"DELETE FROM sections WHERE id IN ({placeholders})"
                        c.execute(query, sectionsIDs)

                        query = f"DELETE FROM chapters_sections_links WHERE sectionID IN ({placeholders})"
                        c.execute(query, sectionsIDs)

                    query = "DELETE FROM chapters WHERE id=(?)"
                    c.execute(query, (chapterID,))
                    conn.commit()
            return True
        except Exception:
            tools.log_traceback()
            return False

    def chapter_links_are_deleted(self, projectID:str, chapterID:int) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        try:
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "DELETE FROM chapters_sections_links WHERE chapterID=(?)"
                    c.execute(query, (chapterID,))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def project_is_deleted(self, projectID:str) -> bool:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            if not Path(projectDirectory).exists():
                raise FileNotFoundError(f'Path {projectDirectory} does not exist.')
            shutil.rmtree(projectDirectory)
            with contextlib.closing(sqlite3.connect(APPDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "DELETE FROM projects where uuid=(?)"
                    c.execute(query, (projectID, ))
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def chapter_sections_links_are_created(self, projectID:str, chapterID:int, sectionsIDs:list) -> bool:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            entries = [(chapterID, sID) for sID in sectionsIDs]
            with contextlib.closing(sqlite3.connect(projectDB)) as conn:
                with contextlib.closing(conn.cursor()) as c:
                    query = "INSERT INTO chapters_sections_links (chapterID,sectionID) VALUES (?,?)"
                    c.executemany(query, entries)
                    conn.commit()
        except Exception:
            tools.log_traceback()
            return False
        return True

    def get_data_for_omilayer(self, projectID:str, DBpath:str, layerName:str, nrows:str) -> tuple:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            omi = Omilayers(str(Path(projectDirectory).joinpath(DBpath)))
            if nrows == "all":
                df = omi.run(f"SELECT * FROM {layerName}", fetchdf=True)
            else:
                df = omi.run(f"SELECT * FROM {layerName} LIMIT {int(nrows)}", fetchdf=True)
            tablesInfo = omi.run(f"SELECT * FROM tables_info WHERE name='{layerName}'", fetchdf=True)

            layerInfo = tablesInfo['info'].values[0]

            jsonCols = []
            for idx,col in enumerate(df.columns, start=1):
                jsonCols.append({"name":f"col{idx}", "field":f"col{idx}", "align":"left", "label":col, "sortable": True})
            jsonRows = []
            for row in df.values.tolist():
                jsonRows.append({f"col{idx}":val for idx,val in enumerate(row, start=1)})
            return (jsonCols, jsonRows, layerInfo)
        except Exception:
            tools.log_traceback()
            return ([], [], "Cannot load layer data.")

    def get_data_from_omilayer_for_plotting(self, projectID:str, DBpath:str, query:str) -> pd.DataFrame:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            omi = Omilayers(str(Path(projectDirectory).joinpath(DBpath)))
            df = omi.run(query, fetchdf=True)
        except Exception:
            tools.log_traceback()
            return pd.DataFrame()
        return df

    def get_section_directory_tree(self, projectID:str, sectionID:int) -> list:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        def build_tree(_path):
            tree = []
            for entry in sorted(_path.iterdir()):
                node = {'label': entry.name}
                if entry.is_dir():
                    node['children'] = build_tree(entry)
                tree.append(node)
            return tree
        try:
            sectionPath = Path(projectDirectory).joinpath(f"sections/{sectionID}")
            if sectionPath.exists():
                sectionPathTree = build_tree(sectionPath)
            else:
                sectionPathTree = []
        except Exception:
            tools.log_traceback()
            return []
        return sectionPathTree

    def omilayer_exists(self, projectID:str, dbPath:str, layerName:str) -> bool:
        projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
        dbFullPath = Path(projectDirectory).joinpath(dbPath)
        omi = Omilayers(str(dbFullPath))
        if omi._dbutils._table_exists(layerName):
            return True
        return False

    def new_omilayer_is_created(self, projectID:str, dbPath:str, layerName:str, layerDescription:str, columns:list) -> bool:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            dbFullPath = Path(projectDirectory).joinpath(dbPath)

            dtypesMapper = {"TEXT":"VARCHAR", "FLOAT":"DOUBLE", "INTEGER":"BIGINT"}
            queryCols = [f"{col['name']} {dtypesMapper[col['datatype']]}" for col in columns]
            with duckdb.connect(str(dbFullPath)) as con:
                query = f"CREATE TABLE {layerName} ({','.join(queryCols)})"
                con.execute(query)

                query = "INSERT INTO tables_info (name) VALUES (?)"
                con.execute(query, [layerName])

                query = "UPDATE tables_info SET info = (?) WHERE name = (?)"
                if layerDescription:
                    con.execute(query, [layerDescription, layerName])
                else:
                    con.execute(query, ["No available description.", layerName])
        except Exception:
            tools.log_traceback()
            return False
        return True

    def get_column_names_and_dtypes_for_omilayer(self, projectID:str, dbRelativePath:str, layerName:str) -> dict:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            dbFullPath = Path(projectDirectory).joinpath(dbRelativePath)
            dtypesMapper = {"VARCHAR":"TEXT", "DOUBLE":"FLOAT", "BIGINT":"INTEGER"}
            with duckdb.connect(str(dbFullPath), read_only=True) as con:
                query = f"DESCRIBE {layerName}"
                cols = con.execute(query).fetchdf()
        except Exception:
            tools.log_traceback()
            return {}
        return [{"name":record['column_name'], "dtype":dtypesMapper[record['column_type']], "value":""} for record in cols.to_dict(orient='records')]

    def data_are_inserted_to_omilayer(self, projectID:str, dbRelativePath:str, layerName:str, layerData:list) -> bool:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            dbFullPath = Path(projectDirectory).joinpath(dbRelativePath)
            data = {}
            for item in layerData:
                if item['dtype'] == 'TEXT':
                    data[item['name']] = [item["value"]]
                else:
                    data[item['name']] = [float(item["value"])]
            df = pd.DataFrame(data)
            omi = Omilayers(str(dbFullPath))
            omi.layers[layerName].insert(df)
        except Exception:
            tools.log_traceback()
            return False
        return True

    def file_is_inserted_to_omilayer(self, projectID:str, dbRelativePath:str, layerName:str, uploaded_file, file_extension:str) -> bool:
        try:
            if file_extension == '.xls':
                df = pd.read_excel(uploaded_file, engine='xlrd')
            elif file_extension == '.xlsx':
                df = pd.read_excel(uploaded_file, engine='openpyxl')
            else:
                # sep=None makes pandas to infer the delimiter
                df = pd.read_csv(uploaded_file, sep=None, engine='python')
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            dbFullPath = Path(projectDirectory).joinpath(dbRelativePath)
            omi = Omilayers(str(dbFullPath))
            omi.layers[layerName].insert(df)
        except Exception:
            tools.log_traceback()
            return False
        return True

    def omilayer_description_is_set(self, projectID:str, dbRelativePath:str, layerName:str, layerInfo:str) -> bool:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            dbFullPath = Path(projectDirectory).joinpath(dbRelativePath)
            omi = Omilayers(str(dbFullPath))
            omi.layers[layerName].set_info(layerInfo)
        except Exception:
            tools.log_traceback()
            return False
        return True

    def omilayer_is_deleted(self, projectID:str, dbRelativePath:str, layerName:str) -> bool:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            dbFullPath = Path(projectDirectory).joinpath(dbRelativePath)
            omi = Omilayers(str(dbFullPath))
            omi.layers.drop([layerName])
        except Exception:
            tools.log_traceback()
            return False
        return True


class SectionUtils:

    def __init__(self, ID:int, title:str, tags:str, content:str, sectionDate:str, projectID:str, projectName:str, section_dir_exists:int):
        self.ID = ID
        self.title = title
        self.tags = json.loads(tags)
        self.content = content
        self.sectionDate = sectionDate
        self.projectID = projectID
        self.projectName = projectName
        self.section_dir_exists = section_dir_exists
        tags = tags.strip()

    def render_content_to_html(self):
        renderer = HtmlRenderer(self.projectID, self.projectName, self.content)
        renderedContent = renderer.render_section_content()
        return {
                "projectID"         : self.projectID,
                "projectName"       : tools.get_project_name_based_on_id(self.projectID),
                "ID"                : self.ID,
                "title"             : self.title,
                "tags"              : self.tags,
                "content"           : renderedContent,
                "date"              : self.sectionDate,
                "section_dir_exists": self.section_dir_exists
                }

    def render_content_to_pdf(self):
        renderer = PDFRenderer(self.projectID, self.projectName, self.content)
        renderedContent = renderer.render_section_content()
        return {
                "projectID"         : self.projectID,
                "projectName"       : tools.get_project_name_based_on_id(self.projectID),
                "ID"                : self.ID,
                "title"             : self.title,
                "tags"              : self.tags,
                "content"           : renderedContent,
                "date"              : self.sectionDate,
                "section_dir_exists": self.section_dir_exists
                }


class AiUtils(ProjectsUtils):

    ai_models_fpath = APPDIR.joinpath("ai_models.json")

    def get_chat_sections(self, projectID:str) -> list:
        try:
            projectDirectory, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
            sectionsResults = self.get_sections_based_on_tag(projectID, "ai-chat", descending=True)
            sections = []
            if sectionsResults:
                colsMapper = {
                            "id":"ID",
                            "section":"title",
                            "tags":"tags",
                            "content":"content",
                            "date":"sectionDate"
                                }
                cols = ['id', 'section', 'tags', 'content', 'date']
                for result in sectionsResults:
                    kwargs = {colsMapper[col]:value for col,value in zip(cols,result)}
                    if not kwargs['content'].strip("\n"):
                        kwargs['content'] = "Section content is emtpy."
                    kwargs['projectID'] = projectID
                    kwargs['projectName'] = tools.get_project_name_based_on_id(projectID)
                    kwargs['section_dir_exists'] = int(Path(projectDirectory).joinpath(f"sections/{kwargs['ID']}").exists())
                    section = SectionUtils(**kwargs)
                    sections.append(section.render_content_to_html())
            return sections
        except Exception:
            tools.log_traceback()
            return []

    def _clean_section_content(self, text):
        # Remove \begin and \end blocks along with any curly bracket tags
        text = re.sub(r'\\begin(?:\{[^}]*\})?.*?\\end(?:\{[^}]*\})?', '', text, flags=re.DOTALL)
        # Remove \ais and \aie blocks
        text = re.sub(r'\\ais.*?\\aie', '', text, flags=re.DOTALL)
        # Replace 3 or more consecutive newlines with exactly 2 newlines
        text = re.sub(r'\n{3,}', '\n\n', text)
        # Optional: strip leading and trailing whitespace from the whole text
        return text.strip()

    def _append_section_content(self, projectID:str, section_id:int, expanded_prompt:list):
        try:
            section_content = self.get_section_raw_content(projectID, section_id)
            section_title = self.get_section_title(projectID, section_id)
            section_content = self._clean_section_content(section_content)
            expanded_prompt.append(f"# {section_title}")
            expanded_prompt.append(section_content)
        except Exception:
            raise SyntaxError

    def _get_file_content(self, file_path:Path) -> str:
        file_extension = file_path.suffix
        file_content = None
        if file_extension == '.pdf':
            if extras_installed:
                file_content = pymupdf4llm.to_markdown(file_path)
            else:
                raise ValueError("extras are not installed")
        elif file_extension == '.docx' or file_extension == '.doc':
            if extras_installed:
                doc = Document(file_path)
                full_text = []
                for para in doc.paragraphs:
                    full_text.append(para.text)
                file_content = "\n".join(full_text)
            else:
                raise ValueError("extras are not installed")
        elif tools.is_flat_file(file_path):
            with open(file_path, 'r') as inf:
                file_content = inf.read() 
        return file_content

    def expand_user_prompt(self, projectID:str, user_prompt:str) -> str:
        prompt_lines = user_prompt.split("\n")
        expanded_prompt = []
        for line in prompt_lines:
            line = line.strip()
            if line.startswith("@section"):
                section_id = int(line.split(":", 1)[1].strip())
                self._append_section_content(projectID, section_id, expanded_prompt, is_strict=True)
            elif line.startswith("@chapter"):
                chapter_id = int(line.split(":", 1)[1].strip())
                sections_ids = self.get_sections_IDs_for_chapter(projectID, chapter_id)
                for section_id in sections_ids:
                    self._append_section_content(projectID, section_id, expanded_prompt, is_strict=False)
            elif line.startswith("@notebook"):
                notebook_id = int(line.split(":", 1)[1].strip())
                chapters_ids = self.get_chapters_ids_for_notebook(projectID, notebook_id)
                for chapter_id in chapters_ids:
                    sections_ids = self.get_sections_IDs_for_chapter(projectID, chapter_id)
                    for section_id in sections_ids:
                        self._append_section_content(projectID, section_id, expanded_prompt, is_strict=False)
            elif line.startswith("@file"):
                file_relative_path = line.split(":", 1)[1].strip()
                projectDir, projectDB = tools.get_paths_for_project_dir_and_db(projectID)
                file_full_path = Path(projectDir).joinpath(file_relative_path)

                if not file_full_path.exists():
                    raise ValueError

                file_content = self._get_file_content(file_full_path)
                if file_content is None:
                    raise ValueError

                expanded_prompt.append(file_content)
            else:
                expanded_prompt.append(line)
        return "\n".join(expanded_prompt)

    def create_chat_history(self, projectID:str, new_prompt:str, chatSections:list) -> list:
        chat_history = []
        if chatSections:
            for section in chatSections:
                previous_user_prompt, ai_response = section.split("## Response")

                user_prompt = self.expand_user_prompt(projectID, previous_user_prompt)

                chat_history.extend([
                    {
                        "role": "user",
                        "parts": [{"text": user_prompt}]
                    },
                    {
                        "role": "model",
                        "parts": [{"text": ai_response}]
                    },

                ])
        new_prompt = self.expand_user_prompt(projectID, new_prompt)
        chat_history.append(
            {
                "role": "user",
                "parts": [{"text": new_prompt}]
            }
        )
        return chat_history

